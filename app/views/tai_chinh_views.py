from django.shortcuts import render, redirect, get_object_or_404
from django.urls import reverse
from app.models import *
from app.forms import *
from openpyxl import Workbook
from django.http import HttpResponse
from openpyxl.utils import get_column_letter
from django.contrib.auth.decorators import login_required
from django.db.models import Q, Sum
import urllib
from openpyxl.styles import Alignment
from app.views.utils.block_chain import Blockchain
from django.db.models import ProtectedError


@login_required
def tai_chinh_home(request):
    block_chain = Blockchain()
    invalid_block_count = len(block_chain.is_valid_chain())
    block_count = len(BlockChain.objects.order_by('index'))-1
    
    unpaid_count = TuitionPayment.objects.filter(status=0).count()
    paid_count = TuitionPayment.objects.filter(status=1).count()
    unpaid_amount = TuitionPayment.objects.filter(status=0).aggregate(Sum('payment_amount'))['payment_amount__sum']
    paid_amount = TuitionPayment.objects.filter(status=1).aggregate(Sum('payment_amount'))['payment_amount__sum']
    return render(request, 'tai_chinh/tai_chinh_home.html', {'unpaid_count': unpaid_count,'paid_count': paid_count,'unpaid_amount': unpaid_amount,'paid_amount': paid_amount,'invalid_block_count': invalid_block_count,'block_count': block_count })

@login_required
def StudentPayment(request,id):
    tuition_payment = get_object_or_404(TuitionPayment,id = id)
    return render(request, 'tai_chinh/student_payment/tuition_detail.html', {'tuition_payment': tuition_payment})


@login_required
def InvalidBlockList(request):
    block_chain = Blockchain()
    invalid_blocks = block_chain.is_valid_chain()
    return render(request, 'tai_chinh/block_chain/invalid_block_list.html', {"invalid_blocks": invalid_blocks})

@login_required
def OtherCollectionPlanList(request):
   alert_content = request.GET.get('alert_content')
   preset_list = OtherCollectionPlan.objects.all().order_by('-id')
   return render(request, 'tai_chinh/other_collection_plan_preset/preset_list.html', {'preset_list': preset_list, 'alert_content': alert_content})

@login_required
def CreateOtherCollectionPlan(request):
    study_programs = StudyProgram.objects.all().order_by('-id')
    faculties= Faculty.objects.all().order_by('-id')
    if request.method == 'POST':
        form = OtherCollectionPlanForm(request.POST)
        data_copy = request.POST.copy()
        data_copy['payment_amount'] = data_copy['payment_amount'].replace(",", "") 
        form.data = data_copy
        if form.is_valid():
            form.save()
            alert_content = 'Create'
            url = reverse('preset_list') + f'?alert_content={alert_content}'
            return redirect(url)
    else:
        form = OtherCollectionPlanForm()
    return render(request, 'tai_chinh/other_collection_plan_preset/create_preset.html', {'form': form, 'study_programs': study_programs,'faculties': faculties})

@login_required
def EditOtherCollectionPlan(request, id):
    other_collection_plan = get_object_or_404(OtherCollectionPlan, id=id)
    if request.method == 'POST':
        other_collection_plan.description = request.POST.get('description')
        other_collection_plan.payment_name = request.POST.get('payment_name')
        other_collection_plan.save()
        alert_content = 'Edit'
        url = reverse('preset_list') + f'?alert_content={alert_content}'
        return redirect(url)
    else:
        other_collection_plan = OtherCollectionPlan.objects.get(id=id)
        form = OtherCollectionPlanForm()
    return render(request, 'tai_chinh/other_collection_plan_preset/edit_preset.html', {'form': form,'other_collection_plan': other_collection_plan})

@login_required
def DeleteOtherCollectionPlan(request, id):
    other_collection_plan = get_object_or_404(OtherCollectionPlan, id=id)
    payment_amount = format(other_collection_plan.payment_amount, ",.0f")
    if request.method == 'POST':
        try:
            other_collection_plan.delete()
            alert_content = 'Delete'
            url = reverse('preset_list') + f'?alert_content={alert_content}'
            return redirect(url)
        except ProtectedError:
            return render(request, 'error.html')

    return render(request, 'tai_chinh/other_collection_plan_preset/delete_preset.html', {'other_collection_plan': other_collection_plan,'payment_amount': payment_amount})


@login_required
def SelectClassDropdown(request):
    study_program_list = StudyProgram.objects.all().order_by('-id')
    major_list = Major.objects.all().order_by('-id')
    faculty_list = Faculty.objects.all().order_by('-id')
    enrollment_batch = [1,2,3,4,5,6,7,8]
    class_list = StudentClass.objects.all().order_by('-id')
    enrollment_batch_dict = {}

    for major in major_list:
        enrollment_batch = 0
        id = major.id
        for student_class in class_list:
            if student_class.major == major and student_class.enrollment_batch > enrollment_batch:
                enrollment_batch = student_class.enrollment_batch
        enrollment_batch_dict[major.id] = enrollment_batch

    # if request.method == 'POST':
    #     report_type = request.POST.get('report_type')
    #     if report_type == 'tongHop':
    #         SummaryReport()

    return render(request, 'tai_chinh/student_payment/select_class_dropdown.html', {'study_program_list': study_program_list, 'major_list': major_list, 'faculty_list': faculty_list, 'enrollment_batch': enrollment_batch, 'class_list': class_list, 'enrollment_batch_dict': enrollment_batch_dict})

@login_required
def StudentList(request,id):
   student_class = get_object_or_404(StudentClass, id=id)
   student_list = Student.objects.filter(student_class=id).order_by('-id')
   tuition_list = TuitionPayment.objects.filter(student__student_class = student_class).order_by('-student__student_id')
   preset_list = OtherCollectionPayment.objects.filter(student__student_class = student_class).order_by('-student__student_id')
   alert_content = request.GET.get('alert_content')
   return render(request, 'tai_chinh/student_payment/student_list.html', {'tuition_list': tuition_list,'student_list': student_list,'preset_list': preset_list, 'student_class': student_class,'alert_content': alert_content,'student_class_id':id})

#Batch insert 
@login_required
def BatchCreateOtherCollectionPlan(request,id):
    student_class = get_object_or_404(StudentClass, id=id)
    student_list = Student.objects.filter(student_class__id = id).order_by('-id')
    preset_list= OtherCollectionPlan.objects.filter(Q(faculty=student_class.faculty) | Q(faculty=None)).order_by('-id')
    if request.method == 'POST':
        form = BatchOtherCollectionPaymentForm(request.POST)
        data_copy = request.POST.copy()
        data_copy['payment_amount'] = data_copy['payment_amount'].replace(",", "") 
        form.data = data_copy
        if form.is_valid():
            plan_id = request.POST.get('other_collection_plan')
            other_collection_plan = get_object_or_404(OtherCollectionPlan, id=plan_id)
            end_date  = request.POST.get('end_date')
            detail  = request.POST.get('detail')
            payment_amount = request.POST.get('payment_amount').replace(",", "") 
            for student in student_list:
                payment = OtherCollectionPayment(student=student, other_collection_plan=other_collection_plan, end_date = end_date , detail=detail,
                                                 payment_amount = payment_amount, status = 0 )
                payment.save()
            alert_content = 'Create'
            url = reverse('student_list', args=[id]) + f'?alert_content={alert_content}'
            return redirect(url)
    else:
        form = BatchOtherCollectionPaymentForm()
    return render(request, 'tai_chinh/other_collection_plan/batch_add_plan.html', {'form': form, 'student_list': student_list,'preset_list': preset_list,'student_class_id':id})

@login_required
def AddOtherCollectionPayment(request,id):
    study_programs = StudyProgram.objects.all().order_by('-id')
    faculties= Faculty.objects.all().order_by('-id')
    student_class = get_object_or_404(StudentClass, id=id)
    student_list = Student.objects.filter(student_class__id = id).order_by('-id')
    preset_list= OtherCollectionPlan.objects.filter(Q(faculty=student_class.faculty) | Q(faculty=None)).order_by('-id')
    if request.method == 'POST':
        form = OtherCollectionPaymentForm(request.POST)
        data_copy = request.POST.copy()
        data_copy['payment_amount'] = data_copy['payment_amount'].replace(",", "") 
        form.data = data_copy
        if form.is_valid():
            form.save()
            alert_content = 'Create'
            url = reverse('student_list', args=[id]) + f'?alert_content={alert_content}'
            return redirect(url)
    else:
        form = OtherCollectionPaymentForm()
    return render(request, 'tai_chinh/other_collection_plan/add_plan.html', {'form': form, 'student_list': student_list,'preset_list': preset_list,'student_class_id':id,'faculties':faculties,'study_programs':study_programs})

@login_required
def DeleteOtherCollectionPayment(request, id):
    other_payment = get_object_or_404(OtherCollectionPayment, id=id)
    student_class_id = other_payment.student.student_class.id

    if request.method == 'POST':
        other_payment.delete()
        alert_content = 'Delete'
        url = reverse('student_list', args=[other_payment.student.student_class.id]) + f'?alert_content={alert_content}'
        return redirect(url)

    return render(request, 'tai_chinh/other_collection_plan/delete_plan.html', {'other_payment': other_payment,'student_class_id':student_class_id})

@login_required
def DetailOtherCollectionPayment(request, id):
    other_tuition_payment = get_object_or_404(OtherCollectionPayment, id=id)
    return render(request, 'tai_chinh/other_collection_plan/detail_plan.html', {'other_tuition_payment': other_tuition_payment})


@login_required
def StudentTuitionList(request,id):
   student = get_object_or_404(Student, id=id)
   tuition_list = TuitionPayment.objects.filter(student=student).order_by('-id')
   preset_list = OtherCollectionPayment.objects.filter(student=student).order_by('-id')
   return render(request, 'tai_chinh/student_payment/student_tuition_list.html', {'tuition_list': tuition_list,'preset_list': preset_list,'student': student})


@login_required
def SelectReport(request):
    study_programs = StudyProgram.objects.all().order_by('-id')
    majors = Major.objects.all().order_by('-id')
    faculties = Faculty.objects.all().order_by('-id')
    years = Year.objects.all().order_by('-id')
    class_years = ClassYear.objects.all().order_by('-id')
    student_classes = StudentClass.objects.all().order_by('-id')
    report_types = ["Báo cáo học phí tổng hợp","Báo cáo học phí tổng hợp theo lớp","Báo cáo công nợ học phí sinh viên","Báo cáo công nợ học phí sinh viên theo lớp"]

    if request.method == 'POST':
        report_type = request.POST.get('report_type')
        start_year = request.POST.get('start_year')
        end_year = request.POST.get('end_year')
        if report_type == 'Báo cáo học phí tổng hợp':
            return SummaryReport(start_year,end_year)
        elif report_type == 'Báo cáo học phí tổng hợp theo lớp':
            student_class = request.POST.get('student_class')
            return SummaryReportByClass(student_class,start_year,end_year)
        elif report_type == 'Báo cáo công nợ học phí sinh viên':
            return StudentDebtReport(start_year,end_year)
        elif report_type == 'Báo cáo công nợ học phí sinh viên theo lớp':
            student_class = request.POST.get('student_class')
            return StudentDebtReportByClass(student_class,start_year,end_year)

    return render(request, 'tai_chinh/generate_report/select_report.html', {'study_programs': study_programs, 'majors': majors, 'faculties': faculties, 'class_years': class_years, 'years': years, 'student_classes': student_classes, 'report_types': report_types})

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image
from django.http import HttpResponse
import urllib.parse

def SummaryReport(start_year,end_year):
    faculty_list = Faculty.objects.all().order_by('-id')
    workbook = Workbook()
    worksheet = workbook.active
    selected_year_range =  start_year + "-" + end_year
    start_range = f"{start_year}-{int(start_year) + 1}"
    end_range = f"{int(end_year) - 1}-{end_year}"
    for faculty in faculty_list:
        queryset = TuitionPayment.objects.filter(student__student_class__faculty=faculty).filter(approved_course__registed_course__class_year__class_year__range=(start_range, end_range)).order_by('student__student_id')
        worksheet = workbook.create_sheet(title=faculty.faculty_name)
        column_names = ['Mã sinh viên', 'Họ tên', 'Ngày sinh','Tên lớp','Năm học','Kỳ học','Số tiền phải nộp','Trạng thái']
        for col_num, column_name in enumerate(column_names, 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.value = column_name
            
            # Adjust the column width to auto-stretch
            column_letter = get_column_letter(col_num)
            worksheet.column_dimensions[column_letter].bestFit = True

        for row_num, obj in enumerate(queryset, 2):
            if obj.status == True:
                status = "Đã hoàn thành"
            else:
                status = "Chưa hoàn thành"

            if obj.approved_course.registed_course.semester == True:
                semester = "Học kỳ 1"
            else:
                semester = "Học kỳ 2"
            row = [
                obj.student.student_id,
                obj.student.student_name,
                obj.student.birth_date,
                obj.student.student_class.class_name,
                obj.approved_course.registed_course.class_year.class_year,
                semester,
                format(obj.payment_amount, ","),
                status

            ]
            for col_num, value in enumerate(row, 1):
                cell = worksheet.cell(row=row_num, column=col_num)
                cell.value = value
                
        # Adjust the column widths to auto-stretch after populating the data
        for column_cells in worksheet.columns:
            length = max(len(str(cell.value)) for cell in column_cells)
            column_letter = get_column_letter(column_cells[0].column)
            worksheet.column_dimensions[column_letter].width = length + 2

        worksheet.insert_rows(1)
        worksheet['A1'] = 'Báo cáo tổng hợp học phí khoa ' + faculty.faculty_name + ' - Năm học ' + selected_year_range
        worksheet.merge_cells('A1:H1')
        merged_cell = worksheet['A1']
        merged_cell.alignment = Alignment(horizontal='center')
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    filename = urllib.parse.quote('Báo cáo tổng hợp.xlsx')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    workbook.remove(workbook['Sheet'])
    workbook.save(response)
    return response


def SummaryReportByClass(student_class_id,start_year,end_year):
    selected_class = get_object_or_404(StudentClass, id=student_class_id)
    selected_year_range =  start_year + "-" + end_year
    start_range = f"{start_year}-{int(start_year) + 1}"
    end_range = f"{int(end_year) - 1}-{end_year}"
    class_year_list = ClassYear.objects.filter(class_year__range=(start_range, end_range)).order_by('-id')
    name = 'Báo cáo tổng hợp theo lớp '+ selected_class.class_name + ' - Năm học '+ selected_year_range +'.xlsx'
    filename = urllib.parse.quote(name)
    workbook = Workbook()
    worksheet = workbook.active
    for class_year in class_year_list:
        queryset = TuitionPayment.objects.filter(Q(approved_course__registed_course__class_year=class_year) & Q(approved_course__registed_course__semester=1) & Q(student__student_class=selected_class)).order_by('student__student_id')
        worksheet = workbook.create_sheet(title=class_year.class_year + " - Học kỳ 1")
        column_names = ['Mã sinh viên', 'Họ tên', 'Ngày sinh','Tên lớp','Năm học','Kỳ học','Số tiền phải nộp','Trạng thái']
        for col_num, column_name in enumerate(column_names, 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.value = column_name
            
            # Adjust the column width to auto-stretch
            column_letter = get_column_letter(col_num)
            worksheet.column_dimensions[column_letter].bestFit = True

        for row_num, obj in enumerate(queryset, 2):
            registed_course = obj.approved_course.registed_course
            if obj.status == True:
                status = "Đã hoàn thành"
            else:
                status = "Chưa hoàn thành"

            if registed_course.semester == True:
                semester = "Học kỳ 1"
            else:
                semester = "Học kỳ 2"
            
            row = [
                obj.student.student_id,
                obj.student.student_name,
                obj.student.birth_date,
                obj.student.student_class.class_name,
                registed_course.class_year.class_year,
                semester,
                format(obj.payment_amount, ","),
                status
            ]
            for col_num, value in enumerate(row, 1):
                cell = worksheet.cell(row=row_num, column=col_num)
                cell.value = value
        
        # Adjust the column widths to auto-stretch after populating the data
        for column_cells in worksheet.columns:
            length = max(len(str(cell.value)) for cell in column_cells)
            column_letter = column_cells[0].column_letter
            worksheet.column_dimensions[column_letter].width = length + 2

        worksheet.insert_rows(1)
        worksheet['A1'] = 'Báo cáo tổng hợp học phí lớp ' + selected_class.class_name + " - Năm học "+ class_year.class_year + " - Học kỳ 1"
        worksheet.merge_cells('A1:H1')
        merged_cell = worksheet['A1']
        merged_cell.alignment = Alignment(horizontal='center')

        queryset = TuitionPayment.objects.filter(Q(approved_course__registed_course__class_year=class_year) & Q(approved_course__registed_course__semester=0) & Q(student__student_class=selected_class)).order_by('student__student_id')
        worksheet = workbook.create_sheet(title=class_year.class_year + " - Học kỳ 2")
        column_names = ['Mã sinh viên', 'Họ tên', 'Ngày sinh','Tên lớp','Năm học','Kỳ học','Số tiền phải nộp','Trạng thái']
        for col_num, column_name in enumerate(column_names, 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.value = column_name
            
            # Adjust the column width to auto-stretch
            column_letter = get_column_letter(col_num)
            worksheet.column_dimensions[column_letter].bestFit = True

        for row_num, obj in enumerate(queryset, 2):
            registed_course = obj.approved_course.registed_course
            if obj.status == True:
                status = "Đã hoàn thành"
            else:
                status = "Chưa hoàn thành"

            if registed_course.semester == True:
                semester = "Học kỳ 1"
            else:
                semester = "Học kỳ 2"
            
            row = [
                obj.student.student_id,
                obj.student.student_name,
                obj.student.birth_date,
                obj.student.student_class.class_name,
                registed_course.class_year.class_year,
                semester,
                format(obj.payment_amount, ","),
                status
            ]
            for col_num, value in enumerate(row, 1):
                cell = worksheet.cell(row=row_num, column=col_num)
                cell.value = value
        
        # Adjust the column widths to auto-stretch after populating the data
        for column_cells in worksheet.columns:
            length = max(len(str(cell.value)) for cell in column_cells)
            column_letter = column_cells[0].column_letter
            worksheet.column_dimensions[column_letter].width = length + 2

        worksheet.insert_rows(1)
        worksheet['A1'] = 'Báo cáo tổng hợp học phí lớp ' + selected_class.class_name + " - Năm học "+ class_year.class_year + " - Học kỳ 2"
        worksheet.merge_cells('A1:H1')
        merged_cell = worksheet['A1']
        merged_cell.alignment = Alignment(horizontal='center')

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    workbook.remove(workbook['Sheet'])
    workbook.save(response)
    return response

def StudentDebtReport(start_year,end_year):
    faculty_list = Faculty.objects.all().order_by('-id')
    selected_year_range =  start_year + "-" + end_year
    start_range = f"{start_year}-{int(start_year) + 1}"
    end_range = f"{int(end_year) - 1}-{end_year}"
    workbook = Workbook()
    worksheet = workbook.active
    for faculty in faculty_list:
        queryset = TuitionPayment.objects.filter(student__student_class__faculty=faculty).filter(approved_course__registed_course__class_year__class_year__range=(start_range, end_range)).order_by('student__student_id').filter(status=0)
        aggregate_queryset = queryset.values('student__student_id','student__student_name','student__birth_date','student__student_class__class_name',).annotate(sum_payment_amount=Sum('payment_amount'))
        worksheet = workbook.create_sheet(title=faculty.faculty_name)
        column_names = ['Mã sinh viên', 'Họ tên', 'Ngày sinh','Tên lớp','Số tiền chưa nộp']

        for col_num, column_name in enumerate(column_names, 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.value = column_name
            
            # Adjust the column width to auto-stretch
            column_letter = get_column_letter(col_num)
            worksheet.column_dimensions[column_letter].bestFit = True

        for row_num, obj in enumerate(aggregate_queryset, 2):
            row = [
                obj['student__student_id'],
                obj['student__student_name'],
                obj['student__birth_date'],
                obj['student__student_class__class_name'],
                format(obj['sum_payment_amount'], ","),
            ]
            for col_num, value in enumerate(row, 1):
                cell = worksheet.cell(row=row_num, column=col_num)
                cell.value = value
        
        # Adjust the column widths to auto-stretch after populating the data
        for column_cells in worksheet.columns:
            length = max(len(str(cell.value)) for cell in column_cells)
            column_letter = column_cells[0].column_letter
            worksheet.column_dimensions[column_letter].width = length + 2
        
        worksheet.insert_rows(1)
        worksheet['A1'] = 'Báo cáo tổng hợp công nợ học phí sinh viên faculty ' + faculty.faculty_name + ' - Năm học ' + selected_year_range
        worksheet.merge_cells('A1:E1')
        merged_cell = worksheet['A1']
        merged_cell.alignment = Alignment(horizontal='center')
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    filename = urllib.parse.quote('Báo cáo công nợ sinh viên - Năm học '+ selected_year_range +'.xlsx')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    workbook.remove(workbook['Sheet'])
    workbook.save(response)
    return response

def StudentDebtReportByClass(student_class_id,start_year,end_year):
    selected_year_range =  start_year + "-" + end_year
    start_range = f"{start_year}-{int(start_year) + 1}"
    end_range = f"{int(end_year) - 1}-{end_year}"
    selected_class = get_object_or_404(StudentClass, id=student_class_id)
    class_year_list = ClassYear.objects.filter(class_year__range=(start_range, end_range)).order_by('-id')
    name = 'Báo cáo công nợ sinh viên theo lớp '+ selected_class.class_name + ' - Năm học '+ selected_year_range +'.xlsx'
    filename = urllib.parse.quote(name)
    workbook = Workbook()
    worksheet = workbook.active
    for class_year in class_year_list:
        queryset = TuitionPayment.objects.filter(Q(approved_course__registed_course__class_year=class_year) & Q(approved_course__registed_course__semester=1) & Q(student__student_class=selected_class)).order_by('student__student_id').filter(status=0)
        aggregate_queryset = queryset.values('student__student_id','student__student_name','student__birth_date','student__student_class__class_name',).annotate(sum_payment_amount=Sum('payment_amount'))
        worksheet = workbook.create_sheet(title=class_year.class_year + " - Học kỳ 1")
        column_names = ['Mã sinh viên', 'Họ tên', 'Ngày sinh','Tên lớp','Số tiền chưa nộp']
        for col_num, column_name in enumerate(column_names, 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.value = column_name
            
            # Adjust the column width to auto-stretch
            column_letter = get_column_letter(col_num)
            worksheet.column_dimensions[column_letter].bestFit = True

        for row_num, obj in enumerate(aggregate_queryset, 2):
            row = [
                obj['student__student_id'],
                obj['student__student_name'],
                obj['student__birth_date'],
                obj['student__student_class__class_name'],
                format(obj['sum_payment_amount'], ","),
            ]
            for col_num, value in enumerate(row, 1):
                cell = worksheet.cell(row=row_num, column=col_num)
                cell.value = value
        
        # Adjust the column widths to auto-stretch after populating the data
        for column_cells in worksheet.columns:
            length = max(len(str(cell.value)) for cell in column_cells)
            column_letter = column_cells[0].column_letter
            worksheet.column_dimensions[column_letter].width = length + 2

        worksheet.insert_rows(1)
        worksheet['A1'] = 'Báo cáo công nợ học phí lớp '+ selected_class.class_name + " - Năm học "+ class_year.class_year + " - Học kỳ 1"
        worksheet.merge_cells('A1:E1')
        merged_cell = worksheet['A1']
        merged_cell.alignment = Alignment(horizontal='center')

        queryset = TuitionPayment.objects.filter(Q(approved_course__registed_course__class_year=class_year) & Q(approved_course__registed_course__semester=0) & Q(student__student_class=selected_class)).order_by('student__student_id').filter(status=0)
        aggregate_queryset = queryset.values('student__student_id','student__student_name','student__birth_date','student__student_class__class_name',).annotate(sum_payment_amount=Sum('payment_amount'))
        worksheet = workbook.create_sheet(title=class_year.class_year + " - Học kỳ 2")
        column_names = ['Mã sinh viên', 'Họ tên', 'Ngày sinh','Tên lớp','Số tiền chưa nộp']
        for col_num, column_name in enumerate(column_names, 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.value = column_name
            
            # Adjust the column width to auto-stretch
            column_letter = get_column_letter(col_num)
            worksheet.column_dimensions[column_letter].bestFit = True

        for row_num, obj in enumerate(aggregate_queryset, 2):
            row = [
                obj['student__student_id'],
                obj['student__student_name'],
                obj['student__birth_date'],
                obj['student__student_class__class_name'],
                format(obj['sum_payment_amount'], ","),
            ]
            for col_num, value in enumerate(row, 1):
                cell = worksheet.cell(row=row_num, column=col_num)
                cell.value = value
        
        # Adjust the column widths to auto-stretch after populating the data
        for column_cells in worksheet.columns:
            length = max(len(str(cell.value)) for cell in column_cells)
            column_letter = column_cells[0].column_letter
            worksheet.column_dimensions[column_letter].width = length + 2

        worksheet.insert_rows(1)
        worksheet['A1'] = 'Báo cáo công nợ học phí lớp '+ selected_class.class_name + " - Năm học "+ class_year.class_year + " - Học kỳ 2"
        worksheet.merge_cells('A1:E1')
        merged_cell = worksheet['A1']
        merged_cell.alignment = Alignment(horizontal='center')
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    workbook.remove(workbook['Sheet'])
    workbook.save(response)
    return response

@login_required
def CollectionPlanList(request):
   alert_content = request.GET.get('alert_content')
   plan_list = CollectionPlan.objects.all().order_by('-id')
   return render(request, 'tai_chinh/collection_plan/plan_list.html', {'plan_list': plan_list, 'alert_content': alert_content})

@login_required
def NewCollectionPlan(request):
    years = Year.objects.all().order_by('-id')
    majors = Major.objects.all().order_by('-id')
    study_programs = StudyProgram.objects.all().order_by('-id')
    if request.method == 'POST':
        form = CollectionPlanForm(request.POST)
        data_copy = request.POST.copy()
        data_copy['payment_amount'] = data_copy['payment_amount'].replace(",", "") 
        form.data = data_copy
        if form.is_valid():
            form.save()
            alert_content = 'Create'
            url = reverse('plan_list') + f'?alert_content={alert_content}'
            return redirect(url)
    else:
        form = CollectionPlanForm()
    return render(request, 'tai_chinh/collection_plan/new_plan.html', {'form': form, 'study_programs': study_programs,'years': years,'majors': majors})

@login_required
def EditCollectionPlan(request, id):
    collection_plan = get_object_or_404(CollectionPlan, id=id)
    years = Year.objects.all().order_by('-id')
    majors = Major.objects.all().order_by('-id')
    study_programs = StudyProgram.objects.all().order_by('-id')
    if request.method == 'POST':
        collection_plan.year =  get_object_or_404(Year, id=request.POST.get('year'))
        collection_plan.payment_name = request.POST.get('payment_name')
        collection_plan.major = get_object_or_404(Major, id=request.POST.get('major'))
        collection_plan.payment_amount = request.POST.get('payment_amount').replace(",", "")
        collection_plan.billing_period = request.POST.get('billing_period')
        collection_plan.start_date = request.POST.get('start_date')
        collection_plan.end_date = request.POST.get('end_date')
        collection_plan.description = request.POST.get('description')
        collection_plan.save()
        alert_content = 'Edit'
        url = reverse('plan_list') + f'?alert_content={alert_content}'
        return redirect(url)
    else:
        collection_plan = CollectionPlan.objects.get(id=id)
        form = CollectionPlanForm()
    return render(request, 'tai_chinh/collection_plan/edit_plan.html', {'form': form,'collection_plan': collection_plan,'study_programs': study_programs,'years': years,'majors': majors})

@login_required
def DeleteCollectionPlan(request, id):
    collection_plan = get_object_or_404(CollectionPlan, id=id)
    payment_amount = format(collection_plan.payment_amount, ",.0f")
    if request.method == 'POST':
        collection_plan.delete()
        alert_content = 'Delete'
        url = reverse('plan_list') + f'?alert_content={alert_content}'
        return redirect(url)

    return render(request, 'tai_chinh/collection_plan/delete_plan.html', {'collection_plan': collection_plan, 'payment_amount': payment_amount})

@login_required
def CollectionPlanDetail(request, id):
    collection_plan = get_object_or_404(CollectionPlan, id=id)
    payment_amount = format(collection_plan.payment_amount, ",.0f")
    return render(request, 'tai_chinh/collection_plan/detail_plan.html', {'collection_plan': collection_plan, 'payment_amount': payment_amount})


@login_required
def TuitionPaymentList(request,id):
   student = get_object_or_404(Student, id=id)
   tuition_list = TuitionPayment.objects.filter(student_id = student.student_id, payment_id__icontains="HP").order_by('-id')
   other_tuition_list = TuitionPayment.objects.filter(student_id = student.student_id, payment_id__icontains="K").order_by('-id')
   return render(request, 'tai_chinh/student_payment/student_tuition_list.html', {'tuition_list': tuition_list,'other_tuition_list': other_tuition_list,'student': student})



